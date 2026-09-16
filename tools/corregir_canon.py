#!/usr/bin/env python3
"""Propuesta de planilla canon Piura corregida. NO pisa la original: escribe
output/canon_piura_propuesta_corregida.xlsx con las celdas cambiadas en amarillo y una hoja
"Cambios 16-sep" que lista cada correccion con su evidencia. Fuente de cada correccion: el
nombre real del formulario resuelto por la API de HubSpot (11_formularios / 13_formularios_canon).
"""
import sys, pathlib, shutil
sys.path.insert(0, str(pathlib.Path(__file__).parent))
from lib import hs
import openpyxl
from openpyxl.styles import PatternFill, Font
from openpyxl.comments import Comment

forms = hs.leer("11_formularios.json")["formularios"]
origen = hs.RAIZ / "input" / "canon_piura.xlsx"
destino = hs.OUT / "canon_piura_propuesta_corregida.xlsx"
shutil.copy(origen, destino)
wb = openpyxl.load_workbook(destino)
ws = wb["Formularios de Google (Piura)"]
AMARILLO = PatternFill("solid", fgColor="FFFF00")
cambios = []

def poner(fila, col, valor, motivo):
    celda = ws.cell(row=fila, column=col)
    antes = celda.value
    celda.value = valor
    celda.fill = AMARILLO
    celda.comment = Comment(f"16-sep-2026 · {motivo}", "auditoria 5minutos")
    cambios.append((fila, ws.cell(row=1, column=col).value, str(antes)[:60] if antes is not None else "", str(valor)[:60], motivo))

def nombre_api(fid):
    return forms[fid]["nombre"]

# Fila 3: URL liderazgo-y-negociacion -> ID 84f82226 y nombre real
poner(3, 3, "Programa de Especialización en Liderazgo y Negociación de Conflictos",
      "los nombres de las filas 3 y 7 estaban cruzados respecto de sus URLs")
poner(3, 4, "84f82226-d1c2-43aa-b4e9-da20f43c0548",
      f"el sitio sirve 84f82226 = '{nombre_api('84f82226-d1c2-43aa-b4e9-da20f43c0548')}'; 43632e61 es otro formulario del mismo programa, no el de la pagina")
# Fila 7: URL gestion-del-talento-y-negociacion -> ID 0c7b3436 y nombre real
poner(7, 3, "Programa de Especialización en Gestión del Talento y Negociación de Conflictos",
      "los nombres de las filas 3 y 7 estaban cruzados respecto de sus URLs")
poner(7, 4, "0c7b3436-4cc6-4ada-91ea-76b82d1e64ed",
      f"el sitio sirve 0c7b3436 = '{nombre_api('0c7b3436-4cc6-4ada-91ea-76b82d1e64ed')}'; 43632e61 es 'PIURA: PDE en Liderazgo y Negociación de Conflictos', otro programa")
# Fila 30: Digital Business Model
poner(30, 4, "00b12055-381c-43f6-9dc1-45372becccd8",
      f"el sitio sirve 00b12055 = '{nombre_api('00b12055-381c-43f6-9dc1-45372becccd8')}'; b92183d3 es un segundo formulario del mismo curso, sin normalizar. Decidir cual queda y archivar el otro")
# Fila 16 y 35: ID con barra al final
poner(16, 4, "20d1a6ad-9788-4a2a-a6cc-4917bed546e8", "el ID traia una barra al final")
poner(35, 4, "a44b1d7f-6c25-4797-a134-6e4581722480", "el ID traia una barra al final")
poner(35, 5, "https://udeponline.pe/curso-de-negocios-innovadores/", "la URL anterior redirige (301) a esta; la pagina existe y responde 200")
# Fila 19: URL mal escrita
poner(19, 5, "https://udeponline.pe/programa-de-especializacion-en-felicidad-y-desarrollo-organizacional/", "faltaba el guion en 'en-felicidad'; la URL anterior da 404")
# Fila 37: duplicada con ID de Meta
poner(37, 4, None, "era un ID de formulario de Meta (1901549703874911), no de HubSpot; la fila 11 ya tiene esta pagina con su ID correcto (d60e69e9)")
poner(37, 3, "DUPLICADA DE LA FILA 11 — eliminar", "fila duplicada")
# Filas 38-41: diplomados sin URL
urls = {"36b4d47d-ccc6-4c5f-add9-b6477706ad65": "https://udeponline.pe/diplomado-en-marketing-digital-y-analitica/",
        "c86ebf10-b97f-49a4-98ca-c16c183462ad": "https://udeponline.pe/diplomado-en-habilidades-para-la-gestion-de-equipos/",
        "74f59f7e-d2cd-40ad-911b-face84652266": "https://udeponline.pe/diplomado-en-liderazgo-y-gestion-del-talento/",
        "0ce67a0b-b785-4f54-b09f-0016e2f3a4d2": "https://udeponline.pe/diplomado-en-gestion-de-la-felicidad-y-bienestar-organizacional/"}
for fila in range(38, 42):
    fid = str(ws.cell(row=fila, column=4).value or "").strip()
    if fid in urls:
        poner(fila, 5, urls[fid], "URL de la pagina publicada que incrusta este formulario (responde 200)")
# Filas nuevas: Analitica Digital y Diplomado Inteligencia Emocional
nuevas = [("Curso", "Analítica Digital & Growth Marketing", "7a38b319-f279-497a-84c6-e85dad3de2a4",
           "https://udeponline.pe/curso-de-analitica-digital-growth-marketing/"),
          ("Diplomado", "Diplomado en Inteligencia Emocional y Coaching para el Liderazgo en las Organizaciones",
           "bef4306e-9106-4ad8-920f-c01e63c7726e",
           "https://udeponline.pe/diplomado-en-inteligencia-emocional-y-coaching-para-el-liderazgo-en-las-organizaciones/")]
fila = 42
for tipo, nombre, fid, url in nuevas:
    for col, val in ((1, "Piura"), (2, tipo), (3, nombre), (4, fid), (5, url)):
        ws.cell(row=fila, column=col).value = val
        ws.cell(row=fila, column=col).fill = AMARILLO
    cambios.append((fila, "fila nueva", "", f"{nombre[:40]} · {fid[:8]}", f"la pagina existe (200) y no tenia fila; nombre real en HubSpot: '{nombre_api(fid)}'"))
    fila += 1

hoja = wb.create_sheet("Cambios 16-sep")
hoja.append(["Fila", "Columna", "Valor anterior", "Valor propuesto", "Motivo / evidencia"])
for c in hoja[1]:
    c.font = Font(bold=True)
for c in cambios:
    hoja.append(list(c))
hoja.column_dimensions["E"].width = 110
hoja.column_dimensions["C"].width = 40
hoja.column_dimensions["D"].width = 45
wb.save(destino)
print(f"output/{destino.name} · {len(cambios)} cambios propuestos")
for c in cambios:
    print(f"  fila {c[0]:>2} · {c[1]}: {c[3][:50]}")
